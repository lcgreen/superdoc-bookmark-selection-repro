"""Generate fresh synthetic documents. No customer document is read."""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

base = Path(__file__).resolve().parents[1]


def fit(text, length):
    assert len(text) <= length
    return text.ljust(length)


chunks = [
    fit(
        "This paragraph contains synthetic text to reproduce a selection issue in the editor. ", 86
    ),
    "Several references follow below. ",
    fit("The example includes four different anchors. The first is called ", 68),
    "reference A",
    ", next is a ",
    "token B",
    " in a second part. ",
    ", next is a ",
    "reference C..",
    fit(" The highlight must follow each rendered line. The reference is ", 69),
    "reference D..",
    " The last marker is printed as END.",
]
for bookmarked in (False, True):
    doc = Document()
    doc.core_properties.author = "Synthetic reproduction"
    doc.core_properties.title = "Bookmark selection reproduction"
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.85)
    section.top_margin = section.bottom_margin = Inches(0.7)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(11)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for element in doc.styles.element.iter(qn("w:pBdr")):
        element.getparent().remove(element)
    style.paragraph_format.line_spacing = 1.1
    if "CommentReference" not in doc.styles:
        doc.styles.add_style("CommentReference", WD_STYLE_TYPE.CHARACTER)
    doc.styles["CommentReference"].font.superscript = True
    doc.add_paragraph("Bookmark selection reproduction", "Title")
    doc.add_paragraph("Select the long example below including its final punctuation.")
    paragraph = doc.add_paragraph()
    paragraph._p.set(qn("w14:paraId"), "00000005")
    for index, text in enumerate(chunks):
        if index == 2:
            empty_run = paragraph.add_run()
            rstyle = OxmlElement("w:rStyle")
            rstyle.set(qn("w:val"), "CommentReference")
            empty_run._r.get_or_add_rPr().append(rstyle)
        is_reference_start = index in (3, 5, 8, 10)
        is_reference_end = index in (3, 6, 8, 10)
        bookmark_id = 105 + 3 * ([3, 5, 8, 10].index(index) if is_reference_start else 0)
        if bookmarked and is_reference_start:
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(bookmark_id))
            start.set(
                qn("w:name"),
                ["orbdur_A1B2C3D4", "orbdur_B2C3D4E5", "orbdur_C3D4E5F6", "orbdur_D4E5F6A7"][
                    [3, 5, 8, 10].index(index)
                ],
            )
            paragraph._p.append(start)
        paragraph.add_run(text)
        if bookmarked and is_reference_end:
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(105 + 3 * [3, 6, 8, 10].index(index)))
            paragraph._p.append(end)
    doc.add_paragraph("The text and run structure are otherwise identical in both fixtures.")
    destination = base / "public" / "fixtures"
    destination.mkdir(parents=True, exist_ok=True)
    doc.save(destination / ("bookmarks.docx" if bookmarked else "no-bookmarks.docx"))
